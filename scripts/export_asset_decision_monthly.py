# -*- coding: utf-8 -*-
"""导出「月度资产配置决策」表（每月具体推荐了哪些资产）。

数据源：frontend/public/data/decision_history.json（212 个月，2009-01 ~ 最新研判月）
产物（写入 output/tables/）：
  1. asset_decision_monthly.csv        逐月资产配置全量表
  2. 月度资产配置决策_2009-2026.xlsx
     Sheet1 逐月资产配置（212 行，冻结首行+筛选）
     Sheet2 资产切换记录（相邻月份推荐资产发生变化的节点）

用法：项目根目录运行  python scripts/export_asset_decision_monthly.py（需 openpyxl）
数据刷新后重跑即可更新。
"""
import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "frontend" / "public" / "data"
OUT = ROOT / "output" / "tables"

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
NOTE_FONT = Font(color="808080", italic=True, size=9)
TYPE_COLORS = {
    "主流资产": "C6EFCE",
    "质量红利": "FFEB9C",
    "价值红利": "FFC7CE",
    "现金": "D9D9D9",
}


def load(name):
    with open(DATA / name, encoding="utf-8") as f:
        return json.load(f)


def yn(v):
    return "是" if v else "否"


def asset_type_label(meta_types, code):
    for a in meta_types:
        if a["code"] == code:
            return a["name"]
    return code


def decision_rows(items, meta_types):
    """逐月资产配置 → 列表行"""
    rows = []
    for x in items:
        rows.append([
            x["month"],
            x["asset_type"],
            asset_type_label(meta_types, x["asset_type"]),
            x["asset_name"],
            x.get("decision_score"),
            x.get("reason") or "",
            yn(x["branches"].get("actual_growth_available", False)),
            yn(x["branches"].get("expected_growth_available", False)),
            yn(x["branches"].get("roe_available", False)),
            yn(x["branches"].get("roe_crowding_high", False)),
            x.get("counts", {}).get("actual_growth_count"),
            x.get("counts", {}).get("expected_growth_count"),
            x.get("counts", {}).get("roe_asset_count"),
            x.get("counts", {}).get("quality_dividend_count"),
            x.get("counts", {}).get("value_dividend_count"),
            x.get("counts", {}).get("selected_count"),
            x.get("actual_growth_trend"),
            x.get("expected_growth_trend"),
            x.get("roe_trend"),
        ])
    return rows


HEADER = ["月份", "资产类型代码", "资产大类", "推荐资产", "决策得分", "决策依据",
          "实际增速信号可用", "预期增速信号可用", "ROE信号可用", "ROE拥挤偏高",
          "实际增速入选数", "预期增速入选数", "ROE入选数", "质量红利入选数", "价值红利入选数",
          "合计入选数", "实际增速趋势", "预期增速趋势", "ROE趋势"]


def style_sheet(ws, nrow):
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{nrow}"
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col[: min(len(col), 5)]:
            pass
    # 列宽（中文按 2 计）
    for col in ws.iter_cols(min_row=1, max_row=nrow):
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col:
            v = cell.value
            if v is None:
                continue
            est = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
            width = max(width, min(est, 40))
        ws.column_dimensions[letter].width = max(width + 2, 8)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    hist = load("decision_history.json")
    meta = load("meta.json")
    items = hist["items"]  # 212 个月
    meta_types = meta["asset_types"]
    rows = decision_rows(items, meta_types)

    # ---- CSV ----
    csv_path = OUT / "asset_decision_monthly.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in rows:
            w.writerow(r)

    # ---- Excel ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "逐月资产配置"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    n = len(rows) + 1
    # 格式
    for i in range(2, n + 1):
        ws.cell(row=i, column=5).number_format = "0.00"       # 决策得分
        for c in (11, 12, 13, 14, 15, 16):
            ws.cell(row=i, column=c).number_format = "0"       # 入选数
        for c in (17, 18, 19):
            ws.cell(row=i, column=c).number_format = "0.00"    # 趋势
        # 资产大类着色
        label = ws.cell(row=i, column=3).value
        if label in TYPE_COLORS:
            ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=TYPE_COLORS[label])
    style_sheet(ws, n)

    # ---- Sheet2 资产切换记录 ----
    ws2 = wb.create_sheet("资产切换记录")
    ws2.append(["切换月份", "切换前资产(大类)", "切换后资产(大类)", "切换前配置", "切换后配置",
                "切换后决策得分", "切换后决策依据"])
    switch_count = 0
    prev = None
    for x in items:
        if prev is not None and x["asset_name"] != prev["asset_name"]:
            switch_count += 1
            ws2.append([
                x["month"],
                asset_type_label(meta_types, prev["asset_type"]),
                asset_type_label(meta_types, x["asset_type"]),
                prev["asset_name"],
                x["asset_name"],
                x.get("decision_score"),
                x.get("reason") or "",
            ])
            ws2.cell(row=ws2.max_row, column=6).number_format = "0.00"
        prev = x
    ws2.append([])
    ws2.append(["注", f"2009-01 起共发生 {switch_count} 次资产切换（相邻月份推荐资产变化）", "", "", "", "", ""])
    for c in range(1, 8):
        ws2.cell(row=ws2.max_row, column=c).font = NOTE_FONT
    style_sheet(ws2, ws2.max_row)

    xlsx_path = OUT / "月度资产配置决策_2009-2026.xlsx"
    wb.save(xlsx_path)
    print("已生成：")
    print(" ", csv_path)
    print(" ", xlsx_path)
    print(f"覆盖 {items[0]['month']} ~ {items[-1]['month']} 共 {len(items)} 个月，资产切换 {switch_count} 次")


if __name__ == "__main__":
    main()
