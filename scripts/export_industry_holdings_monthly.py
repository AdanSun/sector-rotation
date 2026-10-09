# -*- coding: utf-8 -*-
"""导出「每月具体推荐到哪些行业」的表格（最终配置持仓 + 各策略推荐池）。

数据源（frontend/public/data/）：
  - holdings.json        最终策略月度持仓明细：每月实际配置的具体行业（含权重/次月收益）
  - recommendations.json 各策略每月 Top 推荐池（7 策略 × 前5，非最终持仓）
产物（output/tables/）：
  1. 月度行业配置_2009-2026.xlsx
     Sheet1 每月持仓行业一览（宽表：每月一行，按排名排开具体行业）
     Sheet2 持仓明细（含三级行业/权重/评分/次月已实现收益，可复核）
     Sheet3 各策略月度推荐Top（推荐池口径，区别于最终持仓）
  2. industry_holdings_monthly.csv        持仓明细全量
  3. strategy_recommendations_monthly.csv 策略推荐池全量

用法：项目根目录运行  python scripts/export_industry_holdings_monthly.py（需 openpyxl）
数据刷新后重跑即可更新。
"""
import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "frontend" / "public" / "data"
OUT = ROOT / "output" / "tables"

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
NOTE_FONT = Font(color="808080", italic=True, size=9)
PCT = "0.00%"
NUM3 = "0.000"
NUM2 = "0.00"


def load(name):
    with open(DATA / name, encoding="utf-8") as f:
        return json.load(f)


def style_sheet(ws, nrow=None):
    nrow = nrow or ws.max_row
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{nrow}"
    for col in ws.iter_cols(min_row=1, max_row=min(nrow, 500)):
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col:
            v = cell.value
            if v is None:
                continue
            est = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
            width = max(width, min(est, 42))
        ws.column_dimensions[letter].width = max(width + 2, 8)


def round_v(v, nd=6):
    return round(v, nd) if isinstance(v, (int, float)) else v


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    holdings = load("holdings.json")            # 最终持仓 1443 行
    recs = load("recommendations.json")         # 策略推荐池 6704 行
    meta = load("meta.json")
    type_map = {a["code"]: a["name"] for a in meta["asset_types"]}

    # 按月份分组持仓（按 rank 排序）
    by_month = {}
    for x in holdings:
        by_month.setdefault(x["month"], []).append(x)
    months = sorted(by_month)
    max_n = max(len(v) for v in by_month.values())

    # ---------- CSV ----------
    hdr_hold = ["月份", "资产大类", "推荐资产", "排名", "二级行业(展示)", "三级行业",
                "行业代码(三级)", "权重", "评分", "入选来源数", "次月已实现收益(对数)"]
    with open(OUT / "industry_holdings_monthly.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(hdr_hold)
        for x in holdings:
            w.writerow([
                x["month"], type_map.get(x["asset_type"], x["asset_type"]), x["asset_name"],
                x.get("rank"), x.get("display_industry_name"), x.get("industry_name"),
                x.get("industry_code"), round_v(x.get("weight")), round_v(x.get("score")),
                x.get("source_count"), round_v(x.get("next_month_log_return")),
            ])

    hdr_rec = ["月份", "策略", "排名", "二级行业(展示)", "三级行业", "行业代码(三级)",
               "生命周期", "评分", "资产大类", "次月已实现收益(对数)"]
    with open(OUT / "strategy_recommendations_monthly.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(hdr_rec)
        for x in sorted(recs, key=lambda r: (r["month"], r.get("strategy_name") or "", r.get("rank") or 99)):
            w.writerow([
                x["month"], x.get("strategy_name"), x.get("rank"),
                x.get("display_industry_name"), x.get("industry_name"), x.get("industry_code"),
                x.get("lifecycle_stage_zh"), round_v(x.get("score")),
                type_map.get(x.get("asset_type"), x.get("asset_type")),
                round_v(x.get("next_month_log_return")),
            ])

    # ---------- Excel ----------
    wb = openpyxl.Workbook()

    # Sheet1 每月持仓行业一览（宽表）
    ws = wb.active
    ws.title = "每月持仓行业一览"
    head = ["月份", "资产大类", "推荐资产", "持仓行业数"] + [f"行业{i}(rank{i})" for i in range(1, max_n + 1)]
    ws.append(head)
    for m in months:
        rows = sorted(by_month[m], key=lambda x: x.get("rank") or 99)
        r0 = rows[0]
        cells = [m, type_map.get(r0["asset_type"], r0["asset_type"]), r0["asset_name"], len(rows)]
        for i in range(1, max_n + 1):
            # rank i 不一定连续，按位置取
            if i <= len(rows):
                cells.append(rows[i - 1].get("display_industry_name"))
            else:
                cells.append("")
        ws.append(cells)
    # 空列去掉尾部整列为空的
    last_nonempty_col = max(c.column for row in ws.iter_rows() for c in row if c.value is not None)
    if last_nonempty_col < ws.max_column:
        ws.delete_cols(last_nonempty_col + 1, ws.max_column - last_nonempty_col)
    style_sheet(ws)
    ws.freeze_panes = "D2"

    # Sheet2 持仓明细
    ws = wb.create_sheet("持仓明细")
    ws.append(hdr_hold)
    for x in holdings:
        ws.append([
            x["month"], type_map.get(x["asset_type"], x["asset_type"]), x["asset_name"],
            x.get("rank"), x.get("display_industry_name"), x.get("industry_name"),
            x.get("industry_code"), round_v(x.get("weight")), round_v(x.get("score")),
            x.get("source_count"), round_v(x.get("next_month_log_return")),
        ])
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=8).number_format = PCT
        ws.cell(row=i, column=9).number_format = NUM3
        if ws.cell(row=i, column=11).value is not None:
            ws.cell(row=i, column=11).number_format = PCT
    style_sheet(ws)
    ws.append([])
    ws.append(["注", "本表为最终策略月度持仓（模型最终配置）；权重=组合内权重；次月已实现收益为该行业次月真实对数收益（事后核验用）。", "", "", "", "", "", "", "", "", ""])
    for c in range(1, 12):
        ws.cell(row=ws.max_row, column=c).font = NOTE_FONT

    # Sheet3 各策略月度推荐 Top（推荐池，非最终持仓）
    ws = wb.create_sheet("各策略月度推荐Top")
    ws.append(hdr_rec)
    for x in sorted(recs, key=lambda r: (r["month"], r.get("strategy_name") or "", r.get("rank") or 99)):
        ws.append([
            x["month"], x.get("strategy_name"), x.get("rank"),
            x.get("display_industry_name"), x.get("industry_name"), x.get("industry_code"),
            x.get("lifecycle_stage_zh"), round_v(x.get("score")),
            type_map.get(x.get("asset_type"), x.get("asset_type")),
            round_v(x.get("next_month_log_return")),
        ])
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=8).number_format = NUM3
        if ws.cell(row=i, column=10).value is not None:
            ws.cell(row=i, column=10).number_format = PCT
    style_sheet(ws)
    ws.append([])
    ws.append(["注", "本表为 7 个策略各自的月度 Top 推荐池（口径：推荐候选）；最终实际配置见 Sheet1/Sheet2。", "", "", "", "", "", "", "", ""])
    for c in range(1, 11):
        ws.cell(row=ws.max_row, column=c).font = NOTE_FONT

    xlsx_path = OUT / "月度行业配置_2009-2026.xlsx"
    wb.save(xlsx_path)
    print("已生成：")
    print(" ", xlsx_path)
    print(" ", OUT / "industry_holdings_monthly.csv")
    print(" ", OUT / "strategy_recommendations_monthly.csv")
    print(f"持仓覆盖 {months[0]} ~ {months[-1]} 共 {len(months)} 个月，{len(holdings)} 条持仓；策略推荐池 {len(recs)} 条")


if __name__ == "__main__":
    main()
